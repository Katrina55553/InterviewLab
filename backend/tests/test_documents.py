"""Shared document helpers preserve the platform's extraction and privacy behavior."""

import pytest
from docx import Document

from app.core.documents import extract_text, mask_raw_text, parse_demo_resume


def test_extract_txt_and_docx(tmp_path):
    text = "张三\nPython 后端工程师"
    txt = tmp_path / "resume.txt"
    txt.write_text(text, encoding="utf-8")
    assert extract_text(str(txt)) == text

    docx = tmp_path / "resume.docx"
    document = Document()
    for line in text.splitlines():
        document.add_paragraph(line)
    document.save(docx)
    assert extract_text(str(docx)) == text


def test_mask_sensitive_text_before_storage_and_model_calls():
    text = "13812345678 zhangsan@example.com 110101199001011234"
    masked = mask_raw_text(text)
    assert masked == "138****5678 z**n@example.com 1101**********1234"
    assert mask_raw_text(masked) == masked


@pytest.mark.parametrize("check_digit", ["X", "x", "5"])
def test_mask_id_check_digit_before_phone_matching(check_digit):
    value = "11010520000101001" + check_digit
    masked = mask_raw_text("身份证：" + value)
    assert masked == "身份证：1101**********001" + check_digit
    assert mask_raw_text(masked) == masked


def test_extract_docx_table_only_resume(tmp_path):
    document = Document()
    table = document.add_table(rows=2, cols=2)
    cells = [cell for row in table.rows for cell in row.cells]
    for cell, text in zip(cells, ["姓名", "张三", "技能", "Python FastAPI Redis"], strict=True):
        cell.text = text
    path = tmp_path / "table-resume.docx"
    document.save(path)
    assert extract_text(str(path)) == "姓名\n张三\n技能\nPython FastAPI Redis"


def test_extract_docx_keeps_paragraph_table_and_nested_table_order(tmp_path):
    document = Document()
    document.add_paragraph("个人介绍")
    table = document.add_table(rows=1, cols=2)
    cell = table.cell(0, 0).merge(table.cell(0, 1))
    cell.text = "项目经历"
    cell.add_table(rows=1, cols=1).cell(0, 0).text = "嵌套项目详情"
    cell.add_paragraph("项目成果")
    document.add_paragraph("教育经历")
    path = tmp_path / "mixed-resume.docx"
    document.save(path)
    assert extract_text(str(path)) == "个人介绍\n项目经历\n嵌套项目详情\n项目成果\n教育经历"


def test_demo_parser_preserves_resume_structure():
    result = parse_demo_resume("张三\nPython FastAPI Redis\n\n测试大学 本科")
    assert result["name"] == "张三"
    assert result["skills"] == [{"category": "技术栈", "skills": ["Python", "FastAPI", "Redis"]}]
    assert result["education"][0]["school"] == "测试大学 本科"
    assert result["work_experience"] == []
